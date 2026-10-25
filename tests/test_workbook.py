"""Invariants of the rebuilt workbook (docs/MODEL.md section 7), tested on the product.

Structure tests run on the generated file without a spreadsheet engine.
The parity test (invariant 2) recalculates the workbook with LibreOffice and is skipped when soffice is absent.
"""
from __future__ import annotations

import json
import pathlib
import re
import shutil
import subprocess
import sys

import openpyxl
import pytest

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import build_workbook as bw  # noqa: E402

from sales_capacity_model import ModelInputs, run_scenarios  # noqa: E402

SHEET_REF = re.compile(r"'([^']+)'!")
NAME_REF = re.compile(r"\b(v_[a-z0-9_]+|wi_[a-z]+|m_[a-z_]+|st_[a-z]+)\b")


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    out = tmp_path_factory.mktemp("wb") / "model.xlsx"
    bw.main(out)
    return out


@pytest.fixture(scope="module")
def wb(built):
    return openpyxl.load_workbook(built)


def formulas(ws):
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                yield c


# ---------------------------------------------------------------- invariant 4: inputs flow one way
def test_one_way_flow(wb):
    DE, DL, WI, WL = bw.DE, bw.DL, bw.WI, bw.WL
    allowed = {DE: {DL}, DL: {DE}, WL: set(), WI: {WL}}
    for sheet, may_read in allowed.items():
        for c in formulas(wb[sheet]):
            for ref in SHEET_REF.findall(c.value):
                assert ref in may_read, f"{sheet}!{c.coordinate} reads {ref}: {c.value[:60]}"
    # the data logic sheet reads the client sheet only at its input cells (Your value = E, Source = F)
    for c in formulas(wb[DL]):
        for m in re.finditer(r"'" + re.escape(DE) + r"'!\$?([A-Z]+)\$?\d+", c.value):
            assert m.group(1) in ("E", "F"), f"{DL}!{c.coordinate} reads a non-input cell on {DE}: {c.value[:60]}"
    # the what-if logic reads the what-if client sheet only through its choice names
    for c in formulas(wb[WL]):
        for nm in NAME_REF.findall(c.value):
            assert not nm.startswith("wi_") or nm in ("wi_select", "wi_win", "wi_cycle", "wi_deal", "wi_target"), nm


# ---------------------------------------------------------------- invariant 5: named interface
def test_every_name_referenced_exists(wb):
    defined = set(wb.defined_names.keys())
    for ws in wb.worksheets:
        for c in formulas(ws):
            for nm in NAME_REF.findall(c.value):
                assert nm in defined, f"{ws.title}!{c.coordinate} references undefined name {nm}"
    for e in bw.E:
        assert "v_" + e["id"] in defined


def test_whatif_logic_reads_register_only_by_name(wb):
    for c in formulas(wb[bw.WL]):
        assert bw.DE not in c.value and bw.DL not in c.value, c.value[:80]


# ---------------------------------------------------------------- invariant 6: rule in words
def test_every_logic_row_has_its_rule_in_words(wb):
    dl = wb[bw.DL]
    for e in bw.E:
        r = bw.DL_ROW[e["id"]]
        assert dl.cell(r, 6).value, f"no rule text for {e['id']} on {bw.DL}"
    wl = wb[bw.WL]
    for row in wl.iter_rows(min_row=bw.WL_HDR_ROW + 1):
        label, words = row[0].value, row[1].value
        has_formula = any(isinstance(c.value, str) and c.value.startswith("=") for c in row[2:8])
        if has_formula:
            assert words and str(words).strip(), f"{bw.WL} row {row[0].row} ({label}) has formulas but no rule in words"


# ---------------------------------------------------------------- invariant 7: register purity
def test_register_holds_no_scenario_fields(wb):
    for e in bw.E:
        assert not re.search(r"scenario|improvement|what.?if|target", e["id"] + " " + e["name"], re.I), e["id"]
    for c in formulas(wb[bw.DL]):
        assert not NAME_REF.findall(c.value) or all(not n.startswith("wi_") for n in NAME_REF.findall(c.value)), \
            f"register logic reads a scenario choice: {c.coordinate} {c.value[:60]}"


# ---------------------------------------------------------------- invariant 3: unit is the salesperson
def test_throughput_unit_is_salesperson(wb):
    wl = wb[bw.WL]
    labels = {wl.cell(r, 1).value for r in range(bw.WL_HDR_ROW + 1, wl.max_row + 1)}
    assert "Salespeople" in labels
    assert "Revenue per salesperson" in labels
    rev = wl.cell(bw.WL_ROW["revenue"], 3).value
    assert "salespeople" in rev or f"C{bw.WL_ROW['salespeople']}" in rev


# ---------------------------------------------------------------- invariant 2: parity with the original
def _recalc(path: pathlib.Path) -> dict:
    script = pathlib.Path("/mnt/skills/public/xlsx/scripts/recalc.py")
    if shutil.which("soffice") is None or not script.exists():
        pytest.skip("LibreOffice recalculation not available")
    out = subprocess.run([sys.executable, str(script), str(path), "120"], capture_output=True, text=True, timeout=300)
    info = json.loads(out.stdout[out.stdout.index("{"):])
    assert info["total_errors"] == 0, info
    return info


def test_parity_with_original_quality_scenarios(built, tmp_path):
    """Original inputs (the sample defaults) -> same revenue and change in operating income
    for close rate, cycle time, deal size and combined, as the as-is oracle computes them."""
    path = tmp_path / "parity.xlsx"
    shutil.copy(built, path)
    _recalc(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    wl = wb[bw.WL]
    oracle = run_scenarios(ModelInputs())
    cols = {"close_rate": "D", "cycle_time": "E", "deal_size": "F", "combined_quality": "G"}
    base_rev = wl[f"C{bw.WL_ROW['revenue']}"].value
    assert base_rev == pytest.approx(oracle["baseline"].revenue)
    for key, col in cols.items():
        assert wl[f"{col}{bw.WL_ROW['revenue']}"].value == pytest.approx(oracle[key].revenue), key
        assert wl[f"{col}{bw.WL_ROW['d_oi']}"].value == pytest.approx(oracle[key].operating_income_change), key
        assert wl[f"{col}{bw.WL_ROW['trueup']}"].value == "Balanced", key
