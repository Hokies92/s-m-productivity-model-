"""Operational reliability of the blank modern workbook (modern implementation, not the original model).

What is tested, each after recalculation in LibreOffice:
  1. Missing values: with the original's illustrative inputs typed in and any one required value removed, every result
     on Performance breakdown and Impact Analysis is either unchanged or shows "Incomplete" -- never a different number.
     Operating income is "Incomplete" wherever it depends on the missing value; unrelated expense lines stay calculated.
  2. Partial component lines: one blank component line (with no direct entry) makes the line "Incomplete", not a smaller total.
  3. Zeros: an entered 0 is a value; nothing becomes "Incomplete".
  4. Total sales overhead and marketing (original F39) is a calculated total only: no entry cell, no "value used".
  5. Every editable cell is a listed input and drives at least one result in the mode it governs.
  6. Protection survives recalculation: sheets protected without password, inputs unlocked, formulas locked.

© 2000–2026 Growth Enablement. All rights reserved.
"""
from __future__ import annotations

import csv
import importlib.util
import json
import pathlib
import shutil

import openpyxl
import pytest

ROOT = pathlib.Path(__file__).resolve().parents[1]
pytestmark = pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice (soffice) not installed")

DE, PB, IA = "1) Data entry", "2) Performance breakdown", "3) Impact Analysis"
PLAIN = {"territories": "E9", "commission": "E28", "capacity": "E31", "cycle": "E33", "ramp": "E35", "close_rate": "E37",
         "cost_of_revenue_pct": "E59", "rnd_pct": "E60", "gna_pct": "E61"}
LINES = {"deal_size": (11, range(12, 14)), "salary": (15, range(16, 18)), "travel": (19, range(20, 27)),
         "sales_overhead": (41, range(42, 47)), "marketing": (48, range(49, 57))}
SIZES = ["E7", "F7", "G7", "H7"]
COLS = ["C", "E", "F", "G", "H", "J", "K"]


def _load(mod, path):
    spec = importlib.util.spec_from_file_location(mod, path)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


@pytest.fixture(scope="module")
def run(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("rel")
    builder = _load("build_asis_modern", ROOT / "tools" / "build_asis_modern.py")
    recalc = _load("recalc", ROOT / "tools" / "recalc.py")
    blank = builder.main(tmp / "blank.xlsx")
    manifest = json.loads(blank.with_suffix(".manifest.json").read_text())
    orig = {}
    with open(ROOT / "source" / "workbook_cells.csv", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            try:
                orig[(r["sheet"].strip(), r["cell"])] = float(r["value"])
            except ValueError:
                pass
    full_inputs = {(m["sheet"], m["cell"]): orig[(m["orig_sheet"], m["orig_cell"])]
                   for m in manifest["cell_map"] if m["role"] == "input" and (m["orig_sheet"], m["orig_cell"]) in orig}

    variants: dict[str, dict] = {"full": {}}
    for k, cell in PLAIN.items():                                   # 1. a value missing
        variants[f"missing:{k}"] = {(DE, cell): None}
        variants[f"zero:{k}"] = {(DE, cell): 0}
    for k, (row, comps) in LINES.items():
        variants[f"missing:{k}"] = {(DE, f"E{row}"): None, **{(DE, f"E{c}"): None for c in comps}}
        variants[f"zero:{k}"] = {(DE, f"E{row}"): 0}
        for c in comps:                                             # 2. one component line blank, no direct entry
            variants[f"partial:{k}:E{c}"] = {(DE, f"E{row}"): None, (DE, f"E{c}"): None}
            variants[f"zero:{k}:E{c}"] = {(DE, f"E{row}"): None, (DE, f"E{c}"): 0}
    for c in SIZES:
        variants[f"missing:size:{c}"] = {(PB, c): None}
        variants[f"zero:size:{c}"] = {(PB, c): 0}
    variants["direct_overhead"] = {(DE, "E41"): 5_000_000}
    editable = []                                                   # 5. every unlocked cell, in the mode it governs
    wb0 = openpyxl.load_workbook(blank)
    for ws in wb0.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if not c.protection.locked:
                    editable.append((ws.title, c.coordinate))
    for sheet, cell in editable:
        changes = {}
        r = int(cell[1:])
        for k, (row, comps) in LINES.items():
            if sheet == DE and r in comps:
                changes[(DE, f"E{row}")] = None                     # components govern only when the direct entry is blank
        changes[(sheet, cell)] = "PERTURB"
        variants[f"perturb:{sheet}!{cell}"] = changes
        variants[f"perturb_base:{sheet}!{cell}"] = {k: v for k, v in changes.items() if v != "PERTURB"}

    paths = {}
    for name, changes in variants.items():
        wb = openpyxl.load_workbook(blank)
        for (s, c), v in full_inputs.items():
            wb[s][c].value = v
        for (s, c), v in changes.items():
            if v == "PERTURB":
                cur = wb[s][c].value
                if isinstance(cur, (int, float)) and cur:
                    v = cur * 1.07
                else:   # a blank direct entry: enter a figure 1.5x the calculated value it replaces
                    v = 1.5 * {"E11": 75000, "E15": 87750, "E19": 47500, "E41": 10_800_000, "E48": 3_700_000}.get(c, 1)
            wb[s][c].value = v
        p = tmp / (name.replace(":", "_").replace("!", "_").replace(")", "").replace(" ", "_") + ".xlsx")
        wb.save(p)
        paths[name] = p
    paths["blank"] = tmp / "blank_recalc.xlsx"
    shutil.copy(blank, paths["blank"])
    assert recalc.recalc_many(list(paths.values())), "LibreOffice did not recalculate every variant"
    vals = {n: openpyxl.load_workbook(p, data_only=True) for n, p in paths.items()}
    return dict(vals=vals, paths=paths, manifest=manifest, editable=editable)


def results(wbv):
    out = {}
    for s in (PB, IA):
        for r in range(8, 33):
            for col in COLS:
                x = wbv[s][f"{col}{r}"].value
                if x is not None:
                    out[(s, f"{col}{r}")] = x
    return out


def oi(wbv, col="C"):
    return wbv[IA][f"{col}23"].value


def test_full_inputs_complete(run):
    v = run["vals"]["full"]
    assert "complete" in v[DE]["B6"].value and "INCOMPLETE" not in v[IA]["B4"].value
    assert abs(oi(v) - 31_813_750) < 1e-6
    assert not [k for k, x in results(v).items() if not isinstance(x, (int, float))]


@pytest.mark.parametrize("name", [n for n in ["missing:" + k for k in list(PLAIN) + list(LINES)] + [f"missing:size:{c}" for c in SIZES]]
                         + [f"partial:{k}:E{c}" for k, (_, cs) in LINES.items() for c in cs])
def test_missing_value_never_becomes_a_number(run, name):
    full, v = results(run["vals"]["full"]), results(run["vals"][name])
    changed_numbers = [(k, full[k], v[k]) for k in full if v.get(k) != full[k] and v.get(k) != "Incomplete"]
    assert not changed_numbers, f"{name}: results changed to a different number instead of Incomplete: {changed_numbers[:5]}"
    assert any(x == "Incomplete" for x in v.values()), f"{name}: nothing marked Incomplete"
    assert "INCOMPLETE" in run["vals"][name][IA]["B4"].value
    # operating income is never shown as a complete-looking number in a column that depends on the missing value
    for col in COLS:
        dep = any(v.get((IA, f"{col}{r}")) == "Incomplete" for r in range(8, 23))
        if dep:
            assert v[(IA, f"{col}23")] == "Incomplete", f"{name}: operating income {col}23 = {v[(IA, f'{col}23')]}"


@pytest.mark.parametrize("name, still_calculated, incomplete", [
    ("missing:marketing", [8, 12, 13, 18, 19, 21], [14, 15, 20, 22, 23, 27]),
    ("missing:salary", [8, 9, 14, 18, 19, 21], [12, 13, 15, 20, 22, 23, 26, 27]),
    ("missing:cost_of_revenue_pct", [8, 9, 12, 13, 14, 15, 19, 20, 21, 26, 27], [18, 22, 23]),
    ("missing:sales_overhead", [8, 9, 12, 14, 18, 19, 21, 26], [13, 15, 20, 22, 23, 27]),
])
def test_unrelated_lines_stay_on_the_income_statement(run, name, still_calculated, incomplete):
    full, v = run["vals"]["full"][IA], run["vals"][name][IA]
    for col in ("C", "E", "K"):
        for r in still_calculated:
            if full[f"{col}{r}"].value is None:
                continue
            assert v[f"{col}{r}"].value == full[f"{col}{r}"].value, f"{name}: {col}{r}"
        for r in incomplete:
            if full[f"{col}{r}"].value is None:
                continue
            assert v[f"{col}{r}"].value == "Incomplete", f"{name}: {col}{r} = {v[f'{col}{r}'].value}"


DIVISORS = {"zero:cycle": "E33", "zero:deal_size:E12": "E12"}   # the original formulas divide by these: 12 / cycle, revenue / contracts


@pytest.mark.parametrize("name", [n for n in [f"zero:{k}" for k in list(PLAIN) + list(LINES)] + [f"zero:size:{c}" for c in SIZES]
                                  + [f"zero:{k}:E{c}" for k, (_, cs) in LINES.items() for c in cs] if n not in DIVISORS])
def test_entered_zero_is_a_value(run, name):
    v = results(run["vals"][name])
    assert not [k for k, x in v.items() if x == "Incomplete"], f"{name}: an entered 0 was treated as missing"
    assert "complete" in run["vals"][name][DE]["B6"].value


def test_validation_accepts_zero_except_divisors(run):
    """0 is accepted everywhere except the two inputs the original formulas divide by."""
    wb = openpyxl.load_workbook(run["paths"]["blank"])
    for sheet, cell in run["editable"]:
        dv = next(d for d in wb[sheet].data_validations.dataValidation if cell in d.sqref)
        if sheet == DE and cell in DIVISORS.values():
            assert (dv.operator, dv.formula1) in (("greaterThan", "0"), ("greaterThanOrEqual", "1")), f"{cell} accepts 0"
        else:
            accepts_zero = (dv.operator == "greaterThanOrEqual" and dv.formula1 == "0") or (dv.operator == "between" and dv.formula1 == "0")
            assert accepts_zero, f"{sheet}!{cell}: 0 rejected by validation ({dv.operator} {dv.formula1})"


def test_zero_marketing_counts_as_zero(run):
    assert abs(oi(run["vals"]["zero:marketing"]) - (31_813_750 + 3_700_000)) < 1e-6


def test_total_overhead_and_marketing_is_calculated_only(run):
    m, wb = run["manifest"], openpyxl.load_workbook(run["paths"]["blank"])
    ws = wb[DE]
    assert not [i for i in m["inputs"] if i["cell"] in ("E39", "G39")]
    assert ws["E39"].value is None and ws["G39"].value is None
    assert ws["E39"].protection.locked and ws["G39"].protection.locked
    assert not any("E39" in str(dv.sqref) for dv in ws.data_validations.dataValidation)
    assert "Sales_Overhead_And_Marketing" not in wb.defined_names
    assert run["vals"]["full"][DE]["F39"].value == 14_500_000
    assert run["vals"]["direct_overhead"][DE]["F39"].value == 5_000_000 + 3_700_000


def test_every_editable_cell_is_an_input_that_drives_a_result(run):
    declared = {(i["sheet"], i["cell"]) for i in run["manifest"]["inputs"]} | {(s["sheet"], s["cell"]) for s in run["manifest"]["scenario_size_cells"]}
    assert set(run["editable"]) == declared
    inert = []
    for sheet, cell in run["editable"]:
        base = results(run["vals"][f"perturb_base:{sheet}!{cell}"])
        moved = results(run["vals"][f"perturb:{sheet}!{cell}"])
        if base == moved:
            inert.append(f"{sheet}!{cell}")
    assert not inert, f"editable cells that drive nothing: {inert}"


def test_protection_survives_recalculation(run):
    for name in ("full", "missing:marketing", "blank"):
        wb = openpyxl.load_workbook(run["paths"][name])
        for ws in wb.worksheets:
            assert ws.protection.sheet and not ws.protection.password, f"{name}: {ws.title} unprotected"
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        assert c.protection.locked, f"{name}: formula unlocked {ws.title}!{c.coordinate}"
        for sheet, cell in run["editable"]:
            assert not wb[sheet][cell].protection.locked, f"{name}: input locked {sheet}!{cell}"


def test_blank_workbook_shows_incomplete_not_zero(run):
    v = results(run["vals"]["blank"])
    assert v and all(x == "Incomplete" for x in v.values())
    assert "INCOMPLETE" in run["vals"]["blank"][DE]["B6"].value
