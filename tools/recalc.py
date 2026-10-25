"""Recalculate a workbook with LibreOffice (headless) and report formula errors.

Usage:  python tools/recalc.py FILE.xlsx [timeout_seconds]
Prints JSON: {"status": "success"|"error", "total_errors": n, "total_formulas": n, "error_summary": {...}}
The file is recalculated and saved in place. Requires `soffice` on PATH (LibreOffice Calc).
Self-contained so the tests and the rubric run the same way locally and in CI.

© 2000–2026 Growth Enablement. All rights reserved.
"""
from __future__ import annotations

import json
import os
import pathlib
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""
ERRORS = ("#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!", "Err:")


def recalc(path: pathlib.Path, timeout: int = 120) -> dict:
    if shutil.which("soffice") is None:
        return {"status": "error", "message": "soffice not found on PATH"}
    path = path.resolve()
    with tempfile.TemporaryDirectory() as td:
        profile = pathlib.Path(td) / "profile"
        url = profile.as_uri()
        env = dict(os.environ, HOME=td)
        try:
            subprocess.run(["soffice", "--headless", "--terminate_after_init", f"-env:UserInstallation={url}"],
                           capture_output=True, timeout=timeout, env=env)
        except subprocess.TimeoutExpired:
            return {"status": "error", "message": "LibreOffice timed out creating its profile"}
        macro_dir = profile / "user" / "basic" / "Standard"
        if not macro_dir.exists():
            return {"status": "error", "message": "LibreOffice did not create a profile"}
        (macro_dir / "Module1.xba").write_text(MACRO)
        before = path.stat().st_mtime_ns
        cmd = ["soffice", "--headless", "--norestore", f"-env:UserInstallation={url}",
               "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application", str(path)]
        try:
            r = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout, env=env)
        except subprocess.TimeoutExpired:
            return {"status": "error", "message": "LibreOffice timed out"}
        if r.returncode != 0:
            return {"status": "error", "message": (r.stderr or r.stdout)[-400:]}
        if path.stat().st_mtime_ns == before:
            return {"status": "error", "message": "LibreOffice did not save the file"}
    wb = load_workbook(path, data_only=True)
    wbf = load_workbook(path)
    total = sum(1 for ws in wbf.worksheets for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("="))
    summary: dict[str, dict] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERRORS):
                    key = c.value if not c.value.startswith("Err:") else "Err:"
                    e = summary.setdefault(key, {"count": 0, "locations": []})
                    e["count"] += 1
                    if len(e["locations"]) < 20:
                        e["locations"].append(f"{ws.title}!{c.coordinate}")
    n = sum(e["count"] for e in summary.values())
    return {"status": "success" if n == 0 else "error", "total_errors": n, "total_formulas": total, "error_summary": summary}


if __name__ == "__main__":
    p = pathlib.Path(sys.argv[1])
    t = int(sys.argv[2]) if len(sys.argv) > 2 else 120
    out = recalc(p, t)
    print(json.dumps(out, indent=2))
    sys.exit(0 if out.get("status") == "success" else 1)
