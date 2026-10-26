"""Build the as-is model as a blank, general, modern workbook.

Run:  python tools/build_asis_modern.py  ->  workbook/Sales_Capacity_Model.xlsx  +  workbook/Sales_Capacity_Model.manifest.json
Score: python tools/rubric.py workbook/Sales_Capacity_Model.xlsx workbook/Sales_Capacity_Model.manifest.json

What is kept from the original (source/workbook_cells.csv), cell for cell:
  * the five sheets, their rows and columns, every label and description word for word;
  * every formula's arithmetic, unchanged except that it reads the entry sheet by name; the filled workbook
    reproduces the original's values exactly (tools/rubric.py, gate G1–G4).
Input validation (modern, not part of the original): each engine and analysis formula is wrapped in one guard,
    =IF(<every value it reads is a number>, IFERROR(<original arithmetic>,"n/a"), "Incomplete").
    A missing value is never treated as zero: whatever depends on it shows "Incomplete" all the way to
    operating income, while results that do not depend on it are still calculated. An entered 0 is a value.
    A status line on each sheet names the missing inputs.
What is general:
  * no company data anywhere; every input is blank; the scenario sizes default to the original's 10%;
  * values the original typed into the engine (fully loaded salary) or onto the entry sheet beside a
    calculation (deal size) are entries with two ways to fill them: a direct figure, or the calculation
    from their components. The value the model uses is shown in its own column.
  * the three income statement ratios are entries on the entry sheet, not constants on the analysis sheet.
What is modern: named inputs, data validation, locked formulas with unlocked inputs, grouped detail
rows, one typeface, number formats, freeze panes, fit-to-width printing, a rule in words beside
every calculation, attribution to Growth Enablement.

© 2000–2026 Growth Enablement. All rights reserved.
"""
from __future__ import annotations

import csv
import json
import math
import pathlib
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.protection import SheetProtection
from openpyxl.formatting.rule import CellIsRule, FormulaRule

ROOT = pathlib.Path(__file__).resolve().parents[1]
OUT = ROOT / "workbook" / "Sales_Capacity_Model.xlsx"
MANIFEST = OUT.with_suffix(".manifest.json")

O_DE, O_PB, O_IA, O_SD, O_OR = "1) Data entry", "2) Performance breakdown", "3) Impact Analysis", "4) Scenario Diagnostic", "Orientation"
DE, PB, IA, SD, OR = "1) Data entry", "2) Performance breakdown", "3) Impact Analysis", "4) Scenario Diagnostic", "Orientation"
ATTRIB = "© 2000–2026 Growth Enablement. All rights reserved.  Sales Capacity Model: originated 2000, accuracy audit by Ernst & Young 2003, this edition 2026."

# ---------------------------------------------------------------- original cells
ORIG: dict[tuple[str, str], dict] = {}
with open(ROOT / "source" / "workbook_cells.csv", encoding="utf-8") as f:
    for r in csv.DictReader(f):
        ORIG[(r["sheet"].strip(), r["cell"])] = dict(raw=r["formula_or_input"], value=r["value"],
                                                   is_formula=r["is_formula"].strip().lower() == "true")


def otext(sheet, cell) -> str:
    return " ".join(ORIG[(sheet, cell)]["raw"].split())


# ---------------------------------------------------------------- styles
FONT = "Arial"
f_body = Font(name=FONT, size=10)
f_small = Font(name=FONT, size=9, color="404040")
f_bold = Font(name=FONT, size=10, bold=True)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=16, bold=True)
f_h2 = Font(name=FONT, size=12, bold=True)
f_input = Font(name=FONT, size=10, color="0000FF")
fill_input = PatternFill("solid", start_color="FFF2CC")
fill_calc = PatternFill("solid", start_color="F2F2F2")
fill_used = PatternFill("solid", start_color="DDEBF7")
fill_head = PatternFill("solid", start_color="1F3864")
fill_section = PatternFill("solid", start_color="D9E1F2")
fill_total = PatternFill("solid", start_color="E2EFDA")
fill_incomplete = PatternFill("solid", start_color="F8CBAD", end_color="F8CBAD")
f_incomplete = Font(name=FONT, size=10, bold=True, color="9C0006")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
top_rule = Border(top=Side(style="thin", color="000000"), left=thin, right=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
right = Alignment(horizontal="right", vertical="top")
center = Alignment(horizontal="center", vertical="top", wrap_text=True)
UNLOCKED = Protection(locked=False)

FMT = {"currency": '$#,##0;($#,##0);"-"', "pct": "0.0%", "count": "#,##0", "count1": "#,##0.0",
       "months": "0.0", "ratio2": "0.00"}


def rows_for(text: str, chars: float) -> int:
    return max(1, math.ceil(len(text) / chars))


def put(ws, ref, value, font=f_body, fill=None, fmt=None, align=wrap, border=box, unlocked=False):
    c = ws[ref]
    c.value = value
    c.font, c.alignment = font, align
    if border:
        c.border = border
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if unlocked:
        c.protection = UNLOCKED
    return c


def protect(ws):
    ws.protection = SheetProtection(sheet=True, formatCells=False, formatColumns=False, formatRows=False,
                                    insertRows=True, deleteRows=True, sort=False, autoFilter=False,
                                    selectLockedCells=False, selectUnlockedCells=False)


def mark_incomplete(ws, rng, status_cell):
    """Shade cells showing "Incomplete", and the status line while it reports missing values."""
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Incomplete"'], fill=fill_incomplete, font=f_incomplete))
    ws.conditional_formatting.add(status_cell, FormulaRule(formula=[f'ISNUMBER(SEARCH("INCOMPLETE",{status_cell}))'], fill=fill_incomplete, font=f_incomplete))


def status(ws, formula):
    ws.merge_cells("B4:L4")
    put(ws, "B4", formula, f_bold, border=None, align=Alignment(wrap_text=True, vertical="center"))
    ws.row_dimensions[4].height = 28


_TOKEN = re.compile(r"(?:'[^']+'!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?(?![A-Za-z_(])|[A-Za-z_][A-Za-z0-9_]*")
DEFINED: set = set()


def guard(formula: str) -> str:
    """Original formula -> the same arithmetic, calculated only when every value it reads is a number.
    A value that is blank, or itself "Incomplete", makes this result "Incomplete"; nothing is read as zero."""
    body = formula[1:]
    conds = []
    for m in _TOKEN.finditer(body):
        tok = m.group(0)
        if ":" in tok:
            c = f"COUNT({tok})=ROWS({tok})*COLUMNS({tok})"
        elif tok in DEFINED or re.fullmatch(r"(?:'[^']+'!)?\$?[A-Z]{1,3}\$?\d+", tok):
            c = f"ISNUMBER({tok})"
        else:
            continue                                   # a function name
        if c not in conds:
            conds.append(c)
    cond = conds[0] if len(conds) == 1 else "AND(" + ",".join(conds) + ")"
    return f'=IF({cond},IFERROR({body},"n/a"),"Incomplete")'


manifest = dict(sheet_order=[OR, DE, PB, IA, SD], entry_sheet=DE, inputs=[], cell_map=[], scenario_size_cells=[],
                chain_rows=[], detail_rows=[], diagnostic_sheet=SD, no_gridline_sheets=[OR, PB, IA, SD])


def cmap(orig_sheet, orig_cell, sheet, cell, role):
    manifest["cell_map"].append(dict(orig_sheet=orig_sheet, orig_cell=orig_cell, sheet=sheet, cell=cell, role=role))


def name(wb, nm, sheet, ref):
    col = "".join(ch for ch in ref if ch.isalpha())
    row = "".join(ch for ch in ref if ch.isdigit())
    wb.defined_names[nm] = DefinedName(nm, attr_text=f"'{sheet}'!${col}${row}")
    DEFINED.add(nm)


# ---------------------------------------------------------------- 1) Data entry
# row: (kind, unit, name, definition, where, validation)
#   kind: "input" plain entry (original F column input) · "component" (original E column input)
#         · "calc" calculated line with optional direct entry (original F formula) · "ratio" new entry rows
ENTRY = {
    9:  ("input", "count", "Territories", "Sales territories funded last year, filled or open; a territory is a slot in your coverage design, not a person.", "Territory plan or sales operations", "count"),
    11: ("calc", "currency", "Deal_Size", "Average revenue per closed contract. Enter it directly, or leave blank and the model divides revenue by contracts (lines below).", "Finance: revenue ÷ contracts, or CRM closed-won report", "currency"),
    12: ("component", "count", None, "Contracts closed in the year.", "Billing or contract system", "divisor"),
    13: ("component", "currency", None, "Revenue recognized from those contracts.", "Income statement", "currency"),
    15: ("calc", "currency", "Loaded_Salary", "Fully loaded cost of one salesperson per year. Enter it directly, or leave blank and the model computes salary × (1 + load rate).", "Payroll and benefits ledger, quota-carrying roles only", "currency"),
    16: ("component", "currency", None, "Average base salary of a quota-carrying salesperson.", "Payroll register", "currency"),
    17: ("component", "pct", None, "Benefits and payroll taxes as a share of base salary, e.g. 0.23.", "Benefits ledger ÷ base payroll", "pct"),
    19: ("calc", "currency", "Travel_Per_Salesperson", "Travel and entertainment per salesperson per year. Enter it directly, or leave blank and the model builds it from the trip lines below.", "T&E ledger for quota-carrying salespeople ÷ salespeople", "currency"),
    20: ("component", "currency", None, "Air, train or mileage, round trip, one person.", "Expense reports", "currency"),
    21: ("component", "currency", None, "Ground transport per trip.", "Expense reports", "currency"),
    22: ("component", "currency", None, "Meals per day away.", "Expense policy or reports", "currency"),
    23: ("component", "currency", None, "Average room cost per night.", "Expense policy or reports", "currency"),
    24: ("component", "count", None, "Days away per trip.", "Calendars or expense reports", "count"),
    25: ("component", "currency", None, "Client entertainment per trip.", "Expense reports", "currency"),
    26: ("component", "count", None, "Sales trips per salesperson per year.", "Calendars or expense reports", "count"),
    28: ("input", "pct", "Commission_Rate", "Commission as a share of revenue, e.g. 0.08.", "Compensation plan, or commissions paid ÷ revenue", "pct"),
    31: ("input", "count", "Opportunities_Capacity", "The most qualified opportunities one salesperson can hold and advance well at the same time.", "Sales leadership; test against CRM load per salesperson", "count"),
    33: ("input", "months", "Cycle_Months", "Months from qualified opportunity to signature, won deals.", "CRM closed-won report (median)", "months"),
    35: ("input", "months", "Ramp_Months", "Months until a new salesperson reaches average productivity.", "Sales management; productivity by tenure", "months0"),
    37: ("input", "pct", "Close_Rate", "Won ÷ (won + lost) qualified opportunities, e.g. 0.25.", "CRM closed-opportunity report", "pct"),
    39: ("summary", "currency", None, "Total of the two sections below (calculated; not an entry).", "Calculated", None),
    41: ("calc", "currency", "Sales_Overhead", "Non-variable cost of running the sales organization. Enter directly, or leave blank and the model sums the lines below.", "General ledger: sales cost centers", "currency"),
    42: ("component", "currency", None, "Salaries of sales managers who carry no quota.", "Payroll", "currency"),
    43: ("component", "currency", None, "Sales engineers and other non-quota support.", "Payroll", "currency"),
    44: ("component", "currency", None, "Sales operations and administrative staff.", "Payroll", "currency"),
    45: ("component", "currency", None, "All training cost, including travel for training.", "General ledger", "currency"),
    46: ("component", "currency", None, "Non-quota incentives: contests, club, recognition.", "General ledger", "currency"),
    48: ("calc", "currency", "Marketing", "All marketing investment. Enter directly, or leave blank and the model sums the lines below.", "General ledger: marketing cost centers", "currency"),
    49: ("component", "currency", None, "Marketing staff cost.", "Payroll", "currency"),
    50: ("component", "currency", None, "Website build and maintenance.", "General ledger", "currency"),
    51: ("component", "currency", None, "Events, online and in person.", "General ledger", "currency"),
    52: ("component", "currency", None, "Paid media.", "General ledger", "currency"),
    53: ("component", "currency", None, "Direct mail, print and online.", "General ledger", "currency"),
    54: ("component", "currency", None, "Telemarketing and lead generation programs.", "General ledger", "currency"),
    55: ("component", "currency", None, "Brand and public relations.", "General ledger", "currency"),
    56: ("component", "currency", None, "Collateral and sales tools.", "General ledger", "currency"),
    59: ("ratio", "pct", "Cost_Of_Revenue_Pct", "Cost of revenue as a share of revenue, e.g. 0.13.", "Income statement: cost of revenue ÷ revenue", "pct"),
    60: ("ratio", "pct", "RnD_Pct", "Research and development as a share of revenue.", "Income statement: R&D ÷ revenue", "pct"),
    61: ("ratio", "pct", "GnA_Pct", "General and administrative as a share of revenue.", "Income statement: G&A ÷ revenue", "pct"),
}
def _components(rng: str, body: str) -> str:
    """The original calculation from component lines: blank if no line is entered, "Incomplete" if some are
    (a blank line is missing, not zero), otherwise the original arithmetic."""
    return f'=IF(COUNT({rng})=0,"",IF(COUNT({rng})<ROWS({rng}),"Incomplete",IFERROR({body},"n/a")))'


CALC_FORMULA = {   # the original's calculation, kept; the guard is the modern input validation
    11: _components("E12:E13", "E13/E12"),                                      # original typed deal size; this is its stated derivation
    15: _components("E16:E17", "E16+(E16*E17)"),
    19: _components("E20:E26", "(($E$20+$E$21+$E$25)+($E$24*($E$22+$E$23)))*$E$26"),
    39: '=IF(AND(ISNUMBER(G41),ISNUMBER(G48)),G41+G48,"Incomplete")',          # original F39 = F41+F48, a calculated total
    41: _components("E42:E46", "E42+E43+E44+E45+E46"),
    48: _components("E49:E56", "E49+E50+E51+E54+E55+E56+E52+E53"),
}
REQUIRED = [  # (name, label) every value the engine reads from this sheet; used for the status line
    ("Territories", "number of sales territories"), ("Deal_Size", "average revenue per sales transaction"),
    ("Loaded_Salary", "average fully loaded salary"), ("Travel_Per_Salesperson", "average travel costs per salesperson"),
    ("Commission_Rate", "commission"), ("Opportunities_Capacity", "opportunities manageable at one time"),
    ("Cycle_Months", "length of sales cycle"), ("Ramp_Months", "sales ramp up time"), ("Close_Rate", "close rate"),
    ("Sales_Overhead", "sales overhead costs"), ("Marketing", "marketing costs"),
    ("Cost_Of_Revenue_Pct", "cost of revenue ratio"), ("RnD_Pct", "R&D ratio"), ("GnA_Pct", "G&A ratio"),
]
ORIG_INPUT_COL = {9: "F", 11: "F", 28: "F", 31: "F", 33: "F", 35: "F", 37: "F"}   # original typed these in F
DETAIL_GROUPS = [(12, 13), (16, 17), (20, 26), (42, 46), (49, 56)]
UNIT_LABEL = {"currency": "$", "pct": "%", "count": "count", "months": "months"}
ENTRY_RATIO_LABEL = {59: ("Cost of revenue", O_IA, "C30"), 60: ("R&D", O_IA, "C31"), 61: ("G&A (other)", O_IA, "C32")}


def build_entry(ws, wb, dvs):
    ws.sheet_view.showGridLines = True
    for col, w in {"A": 2, "B": 34, "C": 64, "D": 8, "E": 15, "F": 15, "G": 15, "H": 38}.items():
        ws.column_dimensions[col].width = w
    ws.merge_cells("B2:H2")
    put(ws, "B2", otext(O_DE, "B2"), f_title, border=None)
    ws.merge_cells("B3:H3")
    put(ws, "B3", ATTRIB, f_small, border=None)
    ws.merge_cells("B4:H4")
    put(ws, "B4", otext(O_DE, "B4"), f_body, border=None)
    ws.row_dimensions[4].height = 14 * rows_for(otext(O_DE, "B4"), 150)
    ws.merge_cells("B5:H5")
    put(ws, "B5", "How to fill this in: yellow cells are yours to type; grey cells are calculated; the blue column is the value the "
                  "model uses. Where a line can be entered or calculated, a figure typed in the yellow cell is used; otherwise the "
                  "calculation from the lines beneath it is used. A blank is treated as missing, never as zero: enter 0 where an amount "
                  "is truly zero. Results that depend on a missing value show Incomplete.", f_small, border=None)
    ws.row_dimensions[5].height = 38
    for col, h in zip("BCDEFGH", ["Data element", "Description", "Unit", "Entry", "Calculated", "Value used", "Where to find it"]):
        put(ws, f"{col}7", h, f_head, fill_head, align=center)
    ws.freeze_panes = "C8"

    # section rows from the original (B8, B30, B38) and the new ratios section
    for r, text in [(8, otext(O_DE, "B8")), (30, otext(O_DE, "B30")), (38, otext(O_DE, "B38")),
                    (58, "Operating expense ratios (used on 3) Impact Analysis)")]:
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        put(ws, f"B{r}", text, f_bold, fill_section)

    for r, (kind, unit, nm, definition, where, vkind) in ENTRY.items():
        if kind == "summary":                 # original F39: calculated total, nothing to enter
            put(ws, f"B{r}", otext(O_DE, f"B{r}"), f_bold)
            put(ws, f"C{r}", definition, f_small)
            put(ws, f"D{r}", UNIT_LABEL[unit], align=center)
            put(ws, f"E{r}", None)
            put(ws, f"F{r}", CALC_FORMULA[r], f_bold, fill_calc, FMT[unit], align=right)
            put(ws, f"G{r}", None)
            put(ws, f"H{r}", where, f_small)
            cmap(O_DE, f"F{r}", DE, f"F{r}", "formula")
            continue
        if kind == "ratio":
            label, osheet, ocell = ENTRY_RATIO_LABEL[r]
            desc = otext(O_IA, {"C30": "B30", "C31": "B31", "C32": "B32"}[ocell]) + ". " + definition
            put(ws, f"B{r}", label)
            put(ws, f"C{r}", desc, f_small)
        else:
            label = otext(O_DE, f"B{r}") if (O_DE, f"B{r}") in ORIG else None
            desc = otext(O_DE, f"C{r}") if (O_DE, f"C{r}") in ORIG else ""
            if kind == "component":           # the original holds the label in C
                put(ws, f"B{r}", "")
                put(ws, f"C{r}", desc, f_body)
            else:
                put(ws, f"B{r}", label, f_bold if kind == "calc" else f_body)
                put(ws, f"C{r}", desc, f_small)
                ws.row_dimensions[r].height = max(15, 12.5 * rows_for(desc, 95))
        put(ws, f"D{r}", UNIT_LABEL[unit], align=center)
        put(ws, f"E{r}", None, f_input, fill_input, FMT[unit], align=right, unlocked=True)
        dvs[vkind].add(ws[f"E{r}"])
        put(ws, f"H{r}", where, f_small)
        if kind == "calc":
            put(ws, f"F{r}", CALC_FORMULA[r], f_body, fill_calc, FMT[unit], align=right)
            put(ws, f"G{r}", f'=IF(E{r}<>"",E{r},F{r})', f_bold, fill_used, FMT[unit], align=right)
            if ORIG[(O_DE, f"F{r}")]["is_formula"]:
                cmap(O_DE, f"F{r}", DE, f"F{r}", "formula")
            else:                               # the original typed deal size beside its components
                cmap(O_DE, f"F{r}", DE, f"E{r}", "input")
        elif kind in ("input", "ratio"):
            put(ws, f"F{r}", None, f_body, fill_calc)
            put(ws, f"G{r}", f'=IF(E{r}="","",E{r})', f_bold, fill_used, FMT[unit], align=right)
            if kind == "input":
                cmap(O_DE, f"{ORIG_INPUT_COL[r]}{r}", DE, f"E{r}", "input")
            else:
                cmap(*ENTRY_RATIO_LABEL[r][1:], DE, f"E{r}", "input")
        else:
            put(ws, f"F{r}", None, border=box)
            put(ws, f"G{r}", None, border=box)
            cmap(O_DE, f"E{r}", DE, f"E{r}", "input")
        if nm:
            name(wb, nm, DE, f"G{r}")
        manifest["inputs"].append(dict(sheet=DE, cell=f"E{r}", id=(nm or f"entry_{r}").lower(),
                                       label=(label or desc), unit=unit, definition=definition, where=where,
                                       original_text=desc or label))
    ws.merge_cells("B6:H6")
    missing = "&".join(f'IF(ISNUMBER({n}),"","; {lab}")' for n, lab in REQUIRED)
    names = ",".join(n for n, _ in REQUIRED)
    put(ws, "B6", f'=IF(COUNT({names})={len(REQUIRED)},"Status: complete. Every value the model reads has been entered.",'
                  f'"Status: INCOMPLETE. No value for: "&MID({missing},3,2000)&". Results that depend on these show Incomplete.")',
        f_bold, border=None)
    ws.row_dimensions[6].height = 28
    mark_incomplete(ws, "F8:G61", "B6")
    # the original typed the engine salary into the engine (Performance breakdown C11); here it is the direct entry
    cmap(O_PB, "C11", DE, "E15", "input")
    for a, b in DETAIL_GROUPS:
        for r in range(a, b + 1):
            ws.row_dimensions[r].outlineLevel = 1
        manifest["detail_rows"].append(dict(sheet=DE, rows=list(range(a, b + 1))))
    ws.sheet_properties.outlinePr.summaryBelow = False
    protect(ws)


# ---------------------------------------------------------------- 2) Performance breakdown
PB_COLS = ["C", "E", "F", "G", "H", "J", "K"]
PB_RULE = {
    9: "Territories from the entry sheet; the hiring scenario adds the chosen share",
    10: "Deal size from the entry sheet; the deal size scenario adds the chosen share",
    11: "Fully loaded salary from the entry sheet (value used)",
    12: "Travel per salesperson from the entry sheet",
    13: "Commission = revenue per salesperson (row 23) × commission rate",
    14: "Salary + travel + commission",
    16: "Opportunities a salesperson can manage at once, from the entry sheet",
    17: "Sales cycle from the entry sheet; the cycle scenario removes the chosen share",
    18: "Ramp time from the entry sheet",
    19: "Close rate from the entry sheet; the close rate scenario adds the chosen share",
    21: "Opportunities a salesperson can pursue in a year = (12 ÷ cycle) × opportunities managed at once",
    22: "Transactions per salesperson = opportunities pursued × close rate",
    23: "Revenue per salesperson = transactions × deal size",
    25: "Revenue = revenue per salesperson × territories. Hiring: added territories earn (12 − ramp) ÷ 12 of a year. Improve all: the same credit is deducted once",
    26: "Salesperson costs = cost per salesperson × territories",
}
NAME_FOR_ORIG = {"F9": "Territories", "F11": "Deal_Size", "F15": "Loaded_Salary", "F19": "Travel_Per_Salesperson",
                 "$F$28": "Commission_Rate", "F28": "Commission_Rate", "F31": "Opportunities_Capacity", "F33": "Cycle_Months",
                 "F35": "Ramp_Months", "F37": "Close_Rate", "F41": "Sales_Overhead", "$F$41": "Sales_Overhead",
                 "F48": "Marketing", "$F$48": "Marketing"}


def rewrite(formula: str) -> str:
    """Original formula -> the same formula reading the entry sheet by name; then guarded (see guard())."""
    f = formula
    f = re.sub(r"'1\) Data entry '!(\$?F\$?\d+)", lambda m: NAME_FOR_ORIG[m.group(1)], f)
    f = f.replace("'1) Data entry '!", f"'{DE}'!")
    f = f.replace("'2) Performance breakdown'!", f"'{PB}'!").replace("'3) Impact Analysis'!", "")
    return guard(f)


def build_pb(ws, wb, dvs):
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 2, "B": 44, "C": 14, "D": 2, "E": 14, "F": 14, "G": 14, "H": 14, "I": 2, "J": 16, "K": 14, "L": 60}.items():
        ws.column_dimensions[col].width = w
    ws.merge_cells("B2:L2")
    put(ws, "B2", otext(O_PB, "B2"), f_title, border=None)
    ws.merge_cells("B3:L3")
    put(ws, "B3", ATTRIB, f_small, border=None)
    put(ws, "B5", "Scenario", f_head, fill_head)
    put(ws, "C5", "Baseline", f_head, fill_head, align=center)
    ws.merge_cells("E5:H5")
    put(ws, "E5", otext(O_PB, "E5"), f_head, fill_head, align=center)
    ws.merge_cells("J5:K5")
    put(ws, "J5", otext(O_PB, "J5"), f_head, fill_head, align=center)
    put(ws, "L5", "How it is calculated", f_head, fill_head)
    for col in ["E", "F", "G", "H", "J", "K"]:
        put(ws, f"{col}6", otext(O_PB, f"{col}6"), f_bold, fill_section, align=center)
    put(ws, "B6", "", fill=fill_section)
    put(ws, "C6", "", fill=fill_section)
    ws.row_dimensions[6].height = 42
    put(ws, "B7", "Size of improvement (relative, e.g. 0.10 = 10%)", f_bold)
    for col in "EFGH":
        put(ws, f"{col}7", 0.1, f_input, fill_input, FMT["pct"], align=right, unlocked=True)
        dvs["pct"].add(ws[f"{col}7"])
        cmap(O_PB, f"{col}7", PB, f"{col}7", "input")
        manifest["scenario_size_cells"].append(dict(sheet=PB, cell=f"{col}7"))
    put(ws, "J7", "uses the three quality sizes", f_small)
    put(ws, "K7", "uses all four", f_small)
    put(ws, "L7", "Each scenario applies its size to one input and leaves the rest at baseline", f_small)
    ws.freeze_panes = "C8"
    for r in range(8, 29):
        if (O_PB, f"B{r}") not in ORIG:
            continue
        raw = ORIG[(O_PB, f"B{r}")]["raw"]
        is_section = not any((O_PB, f"{c}{r}") in ORIG for c in PB_COLS)
        if r == 28:
            ws.merge_cells("B28:L28")
            put(ws, "B28", otext(O_PB, "B28"), f_small, border=None)
            ws.row_dimensions[28].height = 14 * rows_for(otext(O_PB, "B28"), 190)
            continue
        if is_section:
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=12)
            put(ws, f"B{r}", " ".join(raw.split()), f_bold, fill_section)
            continue
        put(ws, f"B{r}", rewrite(raw) if raw.startswith("=") else " ".join(raw.split()), f_bold if r in (14, 25, 26) else f_body)
        fmt = FMT["currency"] if r in (10, 11, 12, 13, 14, 23, 25, 26) else FMT["pct"] if r == 19 else FMT["count1"] if r in (21, 22) else FMT["months"] if r in (17, 18) else FMT["count"]
        for col in PB_COLS:
            key = (O_PB, f"{col}{r}")
            if key not in ORIG:
                continue
            raw_c = ORIG[key]["raw"]
            if r == 11:
                formula = guard("=Loaded_Salary" if col == "C" else "=C11" if col in ("E", "F", "G", "H", "J") else "=J11")
                role = "formula" if col != "C" else None
            else:
                formula = rewrite(raw_c)
                role = "formula"
            fill = fill_used if r in (25, 26) else fill_calc
            put(ws, f"{col}{r}", formula, f_bold if r in (14, 25, 26) else f_body, fill, fmt, align=right,
                border=top_rule if r in (14, 25) else box)
            if role:
                cmap(O_PB, f"{col}{r}", PB, f"{col}{r}", role)
        put(ws, f"L{r}", PB_RULE[r], f_small)
        ws.row_dimensions[r].height = max(15, 12.5 * rows_for(PB_RULE[r], 70))
        manifest["chain_rows"].append(dict(sheet=PB, row=r, rule_cell=f"L{r}", columns=["E", "F", "G", "H", "J", "K"]))
    status(ws, f"='{DE}'!B6&IF(COUNT(E7:H7)=COLUMNS(E7:H7),\"\",\"  Scenario sizes: INCOMPLETE (a size in row 7 is blank); those scenarios show Incomplete.\")")
    mark_incomplete(ws, "C8:K26", "B4")
    protect(ws)


# ---------------------------------------------------------------- 3) Impact Analysis
IA_RULE = {
    8: "Total revenues from 2) Performance breakdown, row 25",
    9: "Scenario revenue − baseline revenue",
    12: "Scenario salesperson costs − baseline salesperson costs (2) Performance breakdown, row 26)",
    13: "Sales overhead (entry sheet) + salesperson costs",
    14: "Marketing (entry sheet); fixed across scenarios",
    15: "Total sales costs + total marketing costs",
    18: "Revenue × cost of revenue ratio (entry sheet)",
    19: "Revenue × R&D ratio (entry sheet)",
    20: "Total sales costs + total marketing costs",
    21: "Revenue × G&A ratio (entry sheet)",
    22: "Cost of revenue + R&D + sales and marketing + G&A",
    23: "Revenue − total operating expenses",
    26: "(Increase in revenue − increase in sales costs) ÷ baseline revenue",
    27: "(Revenue − sales and marketing costs) ÷ sales and marketing costs",
}
RATIO_NAME = {"$C$30": "Cost_Of_Revenue_Pct", "$C$31": "RnD_Pct", "$C$32": "GnA_Pct"}


def build_ia(ws, wb):
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 2, "B": 36, "C": 15, "D": 2, "E": 15, "F": 15, "G": 15, "H": 15, "I": 2, "J": 17, "K": 15, "L": 56}.items():
        ws.column_dimensions[col].width = w
    ws.merge_cells("B2:L2")
    put(ws, "B2", otext(O_IA, "B2"), f_title, border=None)
    ws.merge_cells("B3:L3")
    put(ws, "B3", ATTRIB, f_small, border=None)
    put(ws, "B5", "", fill=fill_head)
    put(ws, "C5", "", fill=fill_head)
    ws.merge_cells("E5:H5")
    put(ws, "E5", otext(O_IA, "E5"), f_head, fill_head, align=center)
    ws.merge_cells("J5:K5")
    put(ws, "J5", otext(O_IA, "J5"), f_head, fill_head, align=center)
    put(ws, "L5", "How it is calculated", f_head, fill_head)
    put(ws, "B6", "", fill=fill_section)
    for col in ["C", "E", "F", "G", "H", "J", "K"]:
        put(ws, f"{col}6", otext(O_IA, f"{col}6"), f_bold, fill_section, align=center)
    ws.row_dimensions[6].height = 42
    ws.freeze_panes = "C7"
    bold_rows = {8, 15, 22, 23}
    for r in range(7, 28):
        if (O_IA, f"B{r}") not in ORIG:
            continue
        label = otext(O_IA, f"B{r}")
        is_section = not any((O_IA, f"{c}{r}") in ORIG for c in PB_COLS)
        if is_section:
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=12)
            put(ws, f"B{r}", label, f_bold, fill_section)
            continue
        put(ws, f"B{r}", label, f_bold if r in bold_rows else f_body)
        fmt = FMT["pct"] if r == 26 else FMT["ratio2"] if r == 27 else FMT["currency"]
        for col in PB_COLS:
            key = (O_IA, f"{col}{r}")
            if key not in ORIG:
                put(ws, f"{col}{r}", None, border=box)
                continue
            raw = ORIG[key]["raw"]
            f = re.sub(r"\$C\$3[012]", lambda m: RATIO_NAME[m.group(0)], raw)
            formula = rewrite(f)
            fill = fill_total if r == 23 else fill_calc
            put(ws, f"{col}{r}", formula, f_bold if r in bold_rows else f_body, fill, fmt, align=right,
                border=top_rule if r in (15, 22, 23) else box)
            cmap(O_IA, f"{col}{r}", IA, f"{col}{r}", "formula")
        put(ws, f"L{r}", IA_RULE[r], f_small)
        manifest["chain_rows"].append(dict(sheet=IA, row=r, rule_cell=f"L{r}", columns=["E", "F", "G", "H", "J", "K"]))
    # assumptions block: now entered on the entry sheet, shown here
    ws.merge_cells("B29:L29")
    put(ws, "B29", otext(O_IA, "B29") + " — " + otext(O_IA, "E29") + " The three ratios are entered on 1) Data entry (rows 59–61) and shown here.", f_small, border=None)
    ws.row_dimensions[29].height = 14 * rows_for(otext(O_IA, "E29"), 190)
    for r, nm in [(30, "Cost_Of_Revenue_Pct"), (31, "RnD_Pct"), (32, "GnA_Pct")]:
        put(ws, f"B{r}", otext(O_IA, f"B{r}"))
        put(ws, f"C{r}", guard(f"={nm}"), f_body, fill_calc, FMT["pct"], align=right)
    status(ws, f"='{PB}'!B4")
    mark_incomplete(ws, "C7:K32", "B4")
    protect(ws)


# ---------------------------------------------------------------- 4) Scenario Diagnostic
def build_sd(ws):
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 2, "B": 58, "C": 2, "D": 120}.items():
        ws.column_dimensions[col].width = w
    put(ws, "B2", otext(O_SD, "B2"), f_title, border=None)
    ws.merge_cells("B3:D3")
    put(ws, "B3", ATTRIB, f_small, border=None)
    ws.row_dimensions[3].height = 24
    ws.merge_cells("B4:D4")
    put(ws, "B4", otext(O_SD, "B4"), f_body, border=None)
    ws.row_dimensions[4].height = 14 * rows_for(otext(O_SD, "B4"), 200)
    ws.freeze_panes = "A6"
    for (sheet, cell), o in ORIG.items():
        if sheet != O_SD:
            continue
        col, row = cell[0], int(cell[1:])
        if row <= 4:
            continue
        text = " ".join(o["raw"].split())
        is_head = col == "B" and (O_SD, f"D{row}") not in ORIG and (O_SD, f"B{row + 1}") in ORIG and ORIG[(O_SD, f"B{row + 1}")]["raw"].strip() == "Issue"
        if col == "B" and (O_SD, f"D{row}") not in ORIG:
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
            put(ws, f"B{row}", text, f_h2 if is_head else f_bold, fill_section if is_head else None, border=None)
            continue
        if text in ("Issue", "Reason", "Reason / Solution"):
            put(ws, cell, text, f_head, fill_head)
            continue
        put(ws, cell, text, f_bold if col == "B" else f_body)
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 15, 13 * rows_for(text, 150 if col == "D" else 60) + 4)
    protect(ws)


# ---------------------------------------------------------------- Orientation
def build_or(ws):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 150
    put(ws, "B1", otext(O_OR, "B1"), f_title, border=None)
    put(ws, "B2", ATTRIB, f_small, border=None)
    r = 4
    for cell in ("B5", "B6", "B8"):
        text = otext(O_OR, cell)
        put(ws, f"B{r}", text, f_body, border=None)
        ws.row_dimensions[r].height = 14 * rows_for(text, 175)
        r += 2
    put(ws, f"B{r}", "Using this workbook", f_h2, border=None)
    r += 1
    for text in [
        "1)  Data entry: type what you have into the yellow cells. Grey cells are calculated; the blue column is the value the model uses. "
        "Where a line can be entered or calculated, a figure in the yellow cell is used, otherwise the calculation from the lines beneath it.",
        "2)  Performance breakdown: the four scenario sizes are yellow cells at the top. Each scenario changes one input by its size and carries it through the same chain; the two combined columns reuse them.",
        "3)  Impact Analysis: the income statement for the baseline and each scenario, with the percentage net gain and sales and marketing ROI.",
        "4)  Scenario Diagnostic: for each scenario, the common issues that stand between a company and the improvement.",
        "Formulas are locked; inputs are not. Protection has no password: Review > Unprotect Sheet removes it. Every calculated row states its rule in the last column.",
    ]:
        put(ws, f"B{r}", text, f_body, border=None)
        ws.row_dimensions[r].height = 14 * rows_for(text, 175)
        r += 1
    r += 1
    for label, fill, font in [("Yellow: enter", fill_input, f_input), ("Grey: calculated", fill_calc, f_body), ("Blue: value the model uses", fill_used, f_bold)]:
        c = put(ws, f"B{r}", label, font, fill, border=box)
        r += 1
    protect(ws)


# ---------------------------------------------------------------- main
def build():
    wb = Workbook()
    ws_or = wb.active
    ws_or.title = OR
    ws_de = wb.create_sheet(DE)
    ws_pb = wb.create_sheet(PB)
    ws_ia = wb.create_sheet(IA)
    ws_sd = wb.create_sheet(SD)

    def dv_set(ws):
        d = {
            "pct": DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                                  error="Enter a share between 0 and 1 (for example 0.25 for 25%).", errorTitle="Share"),
            "count": DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                                    error="Enter a whole number.", errorTitle="Count"),
            "currency": DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                                       error="Enter an amount in your reporting currency.", errorTitle="Amount"),
            "months": DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True,
                                     error="Enter months, greater than zero.", errorTitle="Months"),
            "months0": DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                                      error="Enter months, zero or more.", errorTitle="Months"),
            "divisor": DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1", allow_blank=True,
                                      error="Enter a whole number of at least 1 (deal size divides revenue by it).", errorTitle="Count"),
        }
        for v in d.values():
            ws.add_data_validation(v)
        return d

    build_or(ws_or)
    build_entry(ws_de, wb, dv_set(ws_de))
    build_pb(ws_pb, wb, dv_set(ws_pb))
    build_ia(ws_ia, wb)
    build_sd(ws_sd)
    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_setup.orientation = "landscape"
    wb.properties.creator = "Growth Enablement"
    wb.properties.title = "Sales Capacity Model"
    return wb


def main(out: pathlib.Path = OUT):
    wb = build()
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    out.with_suffix(".manifest.json").write_text(json.dumps(manifest, indent=1))
    print(out)
    return out


if __name__ == "__main__":
    main()
